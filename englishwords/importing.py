"""Two-step spreadsheet import: bounded parsing, reviewable changes, atomic commit."""
import csv
import hashlib
import io
import json
import zipfile
from collections import Counter, defaultdict

from openpyxl import load_workbook, Workbook
from sqlalchemy import select

from .content import FIELDS, clean, entries, resolved, stamp, bump
from .models import Book, Section, Term, Item, Source

MAX_ROWS = 10000
MAX_COLUMNS = 60
MAX_CELL = 6000
ALIASES = {
    "spelling": ["单词", "英文", "word", "spelling"], "meaning": ["中文释义", "释义", "meaning"],
    "us": ["美音", "美式音标", "phonetic_us", "us"], "uk": ["英音", "英式音标", "uk"],
    "section": ["单元", "章节", "section"], "position": ["顺序", "position"],
    "pos": ["词性", "pos"], "example": ["例句", "example"], "example_zh": ["例句翻译", "example_zh"],
    "book_code": ["书本编号", "book_code"], "book": ["书名", "book"],
    "stage": ["学段", "stage"], "grade": ["年级", "grade"], "publisher": ["出版社", "publisher"],
    "edition": ["教材版本", "edition"], "volume": ["册次", "volume"],
    "source": ["来源", "source"], "notes": ["备注", "notes"],
    "id": ["条目 ID", "条目ID", "id"], "revision": ["修订号", "revision"],
}
LABELS = {k: v[0] for k, v in ALIASES.items()}


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def table(rows):
    result = []
    for number, row in enumerate(rows, 1):
        if number > MAX_ROWS + 1:
            raise ValueError(f"单个工作表最多 {MAX_ROWS} 条记录。")
        if len(row) > MAX_COLUMNS:
            raise ValueError(f"最多 {MAX_COLUMNS} 列。")
        values = []
        for cell in row:
            if hasattr(cell, "data_type"):
                if cell.data_type == "f":
                    raise ValueError(f"第 {number} 行含公式，请转换为文本后导入。")
                cell = cell.value
            value = clean(cell)
            if len(value) > MAX_CELL or "\x00" in value:
                raise ValueError(f"第 {number} 行含过长或不支持的文本。")
            values.append(value)
        if any(values):
            result.append({"line": number, "cells": values})
    if not result:
        raise ValueError("表格为空。")
    headers = result.pop(0)["cells"]
    if len(set(headers)) != len(headers):
        raise ValueError("表头重名，请为每一列填写不同的列名。")
    return {"headers": headers, "rows": result}


def parse_files(uploads):
    files, original = [], io.BytesIO()
    total, row_count = 0, 0
    with zipfile.ZipFile(original, "w", zipfile.ZIP_DEFLATED) as archive:
        for index, upload in enumerate(uploads):
            if index >= 10:
                raise ValueError("一次最多上传 10 个文件。")
            name = upload.filename.replace("\\", "/").split("/")[-1]
            raw = upload.read(10 * 1024 * 1024 + 1)
            total += len(raw)
            if total > 10 * 1024 * 1024:
                raise ValueError("上传文件合计最多 10 MiB。")
            sheets = []
            if name.lower().endswith(".xlsx"):
                with zipfile.ZipFile(io.BytesIO(raw)) as container:
                    if len(container.infolist()) > 2000 or sum(z.file_size for z in container.infolist()) > 50 * 1024 * 1024:
                        raise ValueError("Excel 解压后过大。")
                workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
                if len(workbook.sheetnames) > 20:
                    raise ValueError("工作表数量超过 20。")
                for sheet in workbook:
                    if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                        raise ValueError("工作表列数超过上限，请去除多余格式列。")
                    if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                        raise ValueError("工作表行数超过上限，请拆分文件。")
                    if sheet.max_row == 1 and not any(c.value for c in next(sheet.iter_rows(), ())):
                        continue
                    try:
                        frame = table(sheet.iter_rows())
                    except ValueError as error:
                        if str(error) == "表格为空。":
                            continue
                        raise
                    sheets.append({"name": sheet.title, **frame})
                workbook.close()
            elif name.lower().endswith(".csv"):
                sheets.append({"name": "CSV", **table(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))})
            else:
                raise ValueError("只支持 .xlsx 和 UTF-8 .csv。")
            if not sheets:
                raise ValueError(f"{name} 没有可用工作表。")
            row_count += sum(len(s["rows"]) for s in sheets)
            if row_count > MAX_ROWS:
                raise ValueError("本批工作表合计最多 10,000 条记录，请分批导入。")
            files.append({"name": name, "hash": hashlib.sha256(raw).hexdigest(), "sheets": sheets})
            archive.writestr(f"{index + 1}-{name}", raw)
    if not files:
        raise ValueError("请选择文件。")
    return files, original.getvalue()


def default_mapping(headers):
    return {key: next((str(i) for i, header in enumerate(headers) if header in names), "") for key, names in ALIASES.items()}


def mapped_rows(files, options):
    rows = []
    for i, file in enumerate(files):
        choice = options["files"][i]
        sheet = file["sheets"][int(choice["sheet"])]
        mapping = choice["mapping"]
        if mapping.get("spelling", "") == "":
            raise ValueError(f"{file['name']} 必须映射单词列。")
        for row in sheet["rows"]:
            values = {}
            for key, column in mapping.items():
                if column != "":
                    col = int(column)
                    if col < 0 or col >= len(sheet["headers"]):
                        raise ValueError("列映射无效。")
                    values[key] = row["cells"][col] if col < len(row["cells"]) else ""
            for key in ("book_code", "book", "stage", "grade", "publisher", "edition", "volume", "source"):
                values[key] = values.get(key) or choice.get(key) or ("未核实" if key == "edition" else "")
            values["section"] = values.get("section") or "未分单元"
            values["file"] = file["name"]
            values["line"] = row["line"]
            values["file_hash"] = file["hash"]
            try:
                values["position"] = int(values.get("position") or len(rows) + 1)
                if not 1 <= values["position"] <= 1000000:
                    raise ValueError()
            except ValueError:
                raise ValueError(f"{file['name']} 第 {row['line']} 行：顺序须为 1～1,000,000 的整数。")
            rows.append(values)
    return rows


def make_plan(db, files, options):
    rows = mapped_rows(files, options)
    books = {b.code: b for b in db.scalars(select(Book))}
    existing, by_id = defaultdict(list), {}
    for record in entries(db):
        item, term, section, book = record
        existing[(book.code, section.title, term.spelling)].append(record)
        by_id[item.id] = record
    repeats = Counter((r.get("book_code"), r["section"], r.get("spelling")) for r in rows)
    seen_ids, book_defs, planned = set(), {}, []
    for index, row in enumerate(rows):
        entry = {"index": index, "row": row, "changes": {}, "conflicts": {}, "error": "", "action": "create"}
        code, spelling = row.get("book_code", ""), row.get("spelling", "")
        book = books.get(code)
        if not code or not spelling or (not book and not row.get("book")):
            entry["error"] = "必须填写单词、书本编号，新书还需书名。"
        if len(code) > 100:
            entry["error"] = "书本编号最多 100 个字符。"
        definition = {k: row.get(k, "") for k in ("book", "stage", "grade", "publisher", "edition", "volume")}
        if not book:
            if code in book_defs and book_defs[code] != definition:
                entry["error"] = "同一书本编号的元数据不一致。"
            book_defs[code] = definition
        if book:
            entry["book_revision"] = book.revision
        record = None
        item_id = row.get("id")
        if item_id:
            record = by_id.get(item_id)
            if not record or record[3].code != code:
                entry["error"] = "条目 ID 不存在或不属于此书。"
                record = None
            elif row.get("revision") != str(record[0].revision):
                entry["error"] = "修订号已过期，请导出最新版后再编辑。"
            if item_id in seen_ids:
                entry["error"] = "同一条目 ID 在本批出现多次。"
            seen_ids.add(item_id)
        else:
            candidates = existing.get((code, row["section"], spelling), [])
            if candidates:
                if len(candidates) != 1 or repeats[(code, row["section"], spelling)] != 1:
                    entry["error"] = "同词出现多次，无法自动匹配；请使用带条目 ID 的导出表。"
                else:
                    record = candidates[0]
        if record:
            item, term, section, book = record
            current = resolved(*record)
            entry.update(item_id=item.id, item_revision=item.revision, term_revision=term.revision)
            incoming = {k: row.get(k, "") for k in (*FIELDS, "pos", "notes")}
            if spelling != term.spelling:
                incoming["spelling"] = spelling
            if item_id:
                incoming.update(section=row["section"], position=row["position"])
            for key, value in incoming.items():
                if value != "" and value != current.get(key, ""):
                    target = "changes" if key in ("section", "position") or not current.get(key) else "conflicts"
                    entry[target][key] = {"old": current.get(key, ""), "new": value}
            entry["action"] = "update" if entry["changes"] or entry["conflicts"] else "same"
        entry["missing"] = [LABELS[k] for k in ("meaning", "us") if not row.get(k)]
        planned.append(entry)
    return {"rows": planned, "new_books": book_defs, "errors": sum(bool(p["error"]) for p in planned),
            "creates": sum(p["action"] == "create" for p in planned), "updates": sum(p["action"] == "update" for p in planned),
            "conflicts": sum(bool(p["conflicts"]) for p in planned), "same": sum(p["action"] == "same" for p in planned)}


def apply_plan(db, batch, accepted):
    if batch.state != "preview":
        raise ValueError("该批次不可重复确认。")
    fresh = make_plan(db, batch.files, batch.options)
    if digest(fresh) != digest(batch.plan):
        raise ValueError("预览后词库已改变，请重新预览，避免覆盖新修改。")
    if fresh["errors"]:
        raise ValueError("请先修复导入错误。")
    books = {b.code: b for b in db.scalars(select(Book))}
    terms = {t.spelling: t for t in db.scalars(select(Term))}
    sections = {(s.book_id, s.title): s for s in db.scalars(select(Section))}
    sources = {}
    counts = Counter()
    touched = set()
    for entry in fresh["rows"]:
        row = entry["row"]
        code = row["book_code"]
        if code not in books:
            fields = fresh["new_books"][code]
            book = Book(code=code, title=fields["book"], **{k: v for k, v in fields.items() if k != "book"})
            db.add(book)
            db.flush()
            books[code] = book
        book = books[code]
        source_key = (row["file_hash"], row.get("source", ""))
        if source_key not in sources:
            source = Source(name=row.get("source") or row["file"], file_hash=row["file_hash"], details={"batch_id": batch.id, "filename": row["file"], "redistribution": "未核实"})
            db.add(source)
            db.flush()
            sources[source_key] = source.id
        provenance = {"source_id": sources[source_key], "batch_id": batch.id, "line": row["line"], **stamp(row.get("source") or row["file"])}

        def get_section():
            key = (book.id, row["section"])
            if key not in sections:
                section = Section(book_id=book.id, code="section-" + digest(key)[:16], title=row["section"], position=1 + sum(k[0] == book.id for k in sections))
                db.add(section)
                db.flush()
                sections[key] = section
            return sections[key]

        def get_term(spelling):
            if spelling not in terms:
                term = Term(spelling=spelling, kind="phrase" if " " in spelling else "word")
                db.add(term)
                db.flush()
                terms[spelling] = term
            return terms[spelling]

        if entry["action"] == "create":
            item = Item(section_id=get_section().id, term_id=get_term(row["spelling"]).id, position=row["position"],
                        **{k: row.get(k) or None for k in FIELDS}, pos=row.get("pos", ""), notes=row.get("notes", ""),
                        provenance={k: provenance for k in (*FIELDS, "spelling", "pos", "notes") if row.get(k)})
            db.add(item)
            counts["created"] += 1
            touched.add(book.id)
        else:
            item = db.get(Item, entry["item_id"])
            changes = dict(entry["changes"])
            if str(entry["index"]) in accepted:
                changes.update(entry["conflicts"])
            if not changes:
                counts["unchanged"] += 1
                continue
            provenance_map = dict(item.provenance)
            for key, change in changes.items():
                if key == "section":
                    item.section_id = get_section().id
                elif key == "spelling":
                    item.term_id = get_term(change["new"]).id
                else:
                    setattr(item, key, change["new"])
                provenance_map[key] = provenance
            item.provenance = provenance_map
            item.revision += 1
            item.status = "pending"
            touched.add(book.id)
            counts["updated"] += 1
    for book in books.values():
        if book.id in touched:
            bump(book)
    batch.state = "applied"
    batch.report = dict(counts)
    return dict(counts)


def export_xlsx(records):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "words"
    keys = ["book_code", "book", "stage", "grade", "publisher", "edition", "volume", "section", "position", "spelling", "us", "uk", "meaning", "pos", "example", "example_zh", "notes", "id", "revision"]
    sheet.append([LABELS[k] for k in keys])
    for record in records:
        row = resolved(*record)
        sheet.append([row[k] for k in keys])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 22
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
